from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter


# --- Цветовая схема ---
COLOR_HEADER_BG = "28426A"
COLOR_HEADER_FG = "FFFFFF"
COLOR_CARD_BG = "EAF0F9"
COLOR_CARD_VALUE = "3C78C8"
COLOR_TABLE_HEADER = "28426A"
COLOR_TABLE_ROW_ALT = "F4F7FB"
COLOR_TOTAL_BG = "D6E1F2"
COLOR_BORDER = "C8D2E0"
COLOR_MUTED = "7A869A"


# ============================================================
# ШАБЛОН 1: карточки + таблица по менеджерам (для ОКК)
# ============================================================
def generate_excel_cards(results: dict, output_path: str = "report.xlsx"):
    metrics = results["metrics"]
    managers = results["managers"]
    period = results["period"]
    metric_names = list(metrics.keys())

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    ws.sheet_view.showGridLines = False

    thin = Side(style="thin", color=COLOR_BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def set_cell(cell, value=None, font=None, fill=None, align=None, border_=None):
        if value is not None:
            cell.value = value
        if font:
            cell.font = font
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if align:
            cell.alignment = align
        if border_:
            cell.border = border_

    # --- Заголовок ---
    total_cols = 1 + len(metric_names)
    last_col_letter = get_column_letter(total_cols)

    ws.merge_cells(f"A1:{last_col_letter}1")
    set_cell(
        ws["A1"], "ОТЧЁТ ПО СДЕЛКАМ",
        font=Font(name="Calibri", size=20, bold=True, color=COLOR_HEADER_FG),
        fill=COLOR_HEADER_BG,
        align=Alignment(horizontal="left", vertical="center", indent=1),
    )
    ws.row_dimensions[1].height = 40

    ws.merge_cells(f"A2:{last_col_letter}2")
    set_cell(
        ws["A2"], f"Период: {period}",
        font=Font(name="Calibri", size=12, italic=True, color=COLOR_MUTED),
        align=Alignment(horizontal="left", vertical="center", indent=1),
    )
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 10

    # --- Карточки с ИТОГО ---
    row_label, row_value = 4, 5
    for i, metric_name in enumerate(metric_names):
        col = i + 2
        set_cell(
            ws.cell(row=row_label, column=col), metric_name,
            font=Font(name="Calibri", size=11, bold=True, color=COLOR_HEADER_BG),
            fill=COLOR_CARD_BG,
            align=Alignment(horizontal="center", vertical="center"),
            border_=border,
        )
        set_cell(
            ws.cell(row=row_value, column=col), metrics[metric_name]["ИТОГО"],
            font=Font(name="Calibri", size=22, bold=True, color=COLOR_CARD_VALUE),
            fill=COLOR_CARD_BG,
            align=Alignment(horizontal="center", vertical="center"),
            border_=border,
        )
    ws.row_dimensions[row_label].height = 28
    ws.row_dimensions[row_value].height = 42
    ws.row_dimensions[6].height = 12

    # --- Таблица по менеджерам ---
    header_row = 7
    set_cell(
        ws.cell(row=header_row, column=1), "Менеджер",
        font=Font(name="Calibri", size=12, bold=True, color=COLOR_HEADER_FG),
        fill=COLOR_TABLE_HEADER,
        align=Alignment(horizontal="left", vertical="center", indent=1),
        border_=border,
    )
    for i, metric_name in enumerate(metric_names):
        set_cell(
            ws.cell(row=header_row, column=i + 2), metric_name,
            font=Font(name="Calibri", size=12, bold=True, color=COLOR_HEADER_FG),
            fill=COLOR_TABLE_HEADER,
            align=Alignment(horizontal="center", vertical="center"),
            border_=border,
        )
    ws.row_dimensions[header_row].height = 30

    for idx, manager in enumerate(managers):
        row = header_row + 1 + idx
        row_fill = COLOR_TABLE_ROW_ALT if idx % 2 == 0 else None
        set_cell(
            ws.cell(row=row, column=1), manager,
            font=Font(name="Calibri", size=11, color="333333"),
            fill=row_fill,
            align=Alignment(horizontal="left", vertical="center", indent=1),
            border_=border,
        )
        for i, metric_name in enumerate(metric_names):
            value = metrics[metric_name]["по_менеджерам"].get(manager, 0)
            set_cell(
                ws.cell(row=row, column=i + 2), value,
                font=Font(name="Calibri", size=11, color="333333"),
                fill=row_fill,
                align=Alignment(horizontal="center", vertical="center"),
                border_=border,
            )
        ws.row_dimensions[row].height = 22

    total_row = header_row + 1 + len(managers)
    set_cell(
        ws.cell(row=total_row, column=1), "ИТОГО",
        font=Font(name="Calibri", size=12, bold=True, color=COLOR_HEADER_BG),
        fill=COLOR_TOTAL_BG,
        align=Alignment(horizontal="left", vertical="center", indent=1),
        border_=border,
    )
    for i, metric_name in enumerate(metric_names):
        set_cell(
            ws.cell(row=total_row, column=i + 2), metrics[metric_name]["ИТОГО"],
            font=Font(name="Calibri", size=12, bold=True, color=COLOR_HEADER_BG),
            fill=COLOR_TOTAL_BG,
            align=Alignment(horizontal="center", vertical="center"),
            border_=border,
        )
    ws.row_dimensions[total_row].height = 26

    ws.column_dimensions["A"].width = 30
    for i in range(len(metric_names)):
        col_letter = get_column_letter(i + 2)
        ws.column_dimensions[col_letter].width = 16

    ws.freeze_panes = f"A{header_row + 1}"
    wb.save(output_path)
    print(f"  → Excel сохранён: {output_path}")


# ============================================================
# ШАБЛОН 2: строки (№ | Показатель | Значение | Примечание)
# ============================================================
def generate_excel_rows(
    title: str,
    period: str,
    rows: list[dict],
    output_path: str = "report.xlsx",
):
    """
    rows — список dict'ов вида:
      {"title": "Принято", "value": 12}
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    ws.sheet_view.showGridLines = False

    thin = Side(style="thin", color=COLOR_BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def set_cell(cell, value=None, font=None, fill=None, align=None, border_=None):
        if value is not None:
            cell.value = value
        if font:
            cell.font = font
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if align:
            cell.alignment = align
        if border_:
            cell.border = border_

    # --- Заголовок ---
    ws.merge_cells("A1:D1")
    set_cell(
        ws["A1"], title,
        font=Font(name="Calibri", size=18, bold=True, color=COLOR_HEADER_FG),
        fill=COLOR_HEADER_BG,
        align=Alignment(horizontal="left", vertical="center", indent=1),
    )
    ws.row_dimensions[1].height = 38

    ws.merge_cells("A2:D2")
    set_cell(
        ws["A2"], f"Период: {period}",
        font=Font(name="Calibri", size=11, italic=True, color=COLOR_MUTED),
        align=Alignment(horizontal="left", vertical="center", indent=1),
    )
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 8

    # --- Шапка таблицы ---
    header_row = 4
    headers = ["№", "Показатель", "Значение", "Примечание"]
    aligns = ["center", "left", "center", "left"]
    for i, (h, a) in enumerate(zip(headers, aligns)):
        set_cell(
            ws.cell(row=header_row, column=i + 1), h,
            font=Font(name="Calibri", size=12, bold=True, color=COLOR_HEADER_FG),
            fill=COLOR_TABLE_HEADER,
            align=Alignment(horizontal=a, vertical="center",
                            indent=1 if a == "left" else 0),
            border_=border,
        )
    ws.row_dimensions[header_row].height = 28

        # --- Строки данных ---
    data_idx = 0   # нумерация только для обычных строк (разделители не считаются)
    for idx, row in enumerate(rows):
        r = header_row + 1 + idx

        # === РАЗДЕЛИТЕЛЬ (СДЕЛКИ / ЛИДЫ) ===
        if row.get("header"):
            # Объединяем 4 ячейки под заголовок секции
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            set_cell(
                ws.cell(row=r, column=1),
                f"  {row['title']}",
                font=Font(name="Calibri", size=13, bold=True, color=COLOR_HEADER_FG),
                fill=COLOR_HEADER_BG,
                align=Alignment(horizontal="left", vertical="center", indent=1),
                border_=border,
            )
            ws.row_dimensions[r].height = 30
            continue

        # === ОБЫЧНАЯ СТРОКА ===
        data_idx += 1
        row_fill = COLOR_TABLE_ROW_ALT if data_idx % 2 == 0 else None

        set_cell(
            ws.cell(row=r, column=1), data_idx,
            font=Font(name="Calibri", size=11, color="333333"),
            fill=row_fill,
            align=Alignment(horizontal="center", vertical="center"),
            border_=border,
        )
        set_cell(
            ws.cell(row=r, column=2), row["title"],
            font=Font(name="Calibri", size=11, color="333333"),
            fill=row_fill,
            align=Alignment(horizontal="left", vertical="center", indent=1),
            border_=border,
        )
        set_cell(
            ws.cell(row=r, column=3), row["value"],
            font=Font(name="Calibri", size=11, bold=True, color=COLOR_CARD_VALUE),
            fill=row_fill,
            align=Alignment(horizontal="center", vertical="center"),
            border_=border,
        )
        set_cell(
            ws.cell(row=r, column=4), "",
            font=Font(name="Calibri", size=11, color="333333"),
            fill=row_fill,
            align=Alignment(horizontal="left", vertical="center", indent=1),
            border_=border,
        )
        ws.row_dimensions[r].height = 22

    # --- Ширина колонок ---
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 40

    wb.save(output_path)
    print(f"  → Excel сохранён: {output_path}")